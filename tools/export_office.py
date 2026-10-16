# -*- coding: utf-8 -*-
"""Один файл ОПД_ДЗ_Таблицы.xlsx — все таблицы ДЗ (Эхо-Рыцарь, кузницабагов)."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "ДЗ"
OUT_XLSX = OUT / "ОПД_ДЗ_Таблицы.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
SECTION_FILL = PatternFill("solid", fgColor="2E75B6")
ALT_FILL = PatternFill("solid", fgColor="D6EAF8")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
TITLE_FONT = Font(bold=True, name="Calibri", size=16, color="1F4E79")
SECTION_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=12)
CELL_FONT = Font(name="Calibri", size=10)
BOLD = Font(bold=True, name="Calibri", size=11)
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def style_header_row(ws, row, cols):
    for col in range(1, cols + 1):
        cell = ws.cell(row, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN


def write_table(ws, start_row, headers, rows, alt=True):
    for c, h in enumerate(headers, 1):
        ws.cell(start_row, c, h)
    style_header_row(ws, start_row, len(headers))
    for i, row in enumerate(rows):
        r = start_row + 1 + i
        for c, val in enumerate(row, 1):
            cell = ws.cell(r, c, val)
            cell.font = CELL_FONT
            cell.alignment = WRAP
            cell.border = THIN
            if alt and i % 2 == 1:
                cell.fill = ALT_FILL
        ws.row_dimensions[r].height = max(36, 18 * (1 + max(len(str(v)) for v in row) // 40))
    ws.row_dimensions[start_row].height = 22
    return start_row + 1 + len(rows)


def section_title(ws, row, text, cols=4):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    cell = ws.cell(row, 1, text)
    cell.fill = SECTION_FILL
    cell.font = SECTION_FONT
    cell.alignment = Alignment(vertical="center", horizontal="left")
    for c in range(1, cols + 1):
        ws.cell(row, c).border = THIN
        ws.cell(row, c).fill = SECTION_FILL
    ws.row_dimensions[row].height = 24
    return row + 1


def para(ws, row, text, cols=4, height=48):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    cell = ws.cell(row, 1, text)
    cell.font = CELL_FONT
    cell.alignment = WRAP
    ws.row_dimensions[row].height = height
    return row + 1


def save_xlsx(wb):
    OUT.mkdir(parents=True, exist_ok=True)
    tmp = OUT / "ОПД_ДЗ_Таблицы._tmp.xlsx"
    wb.save(tmp)
    try:
        tmp.replace(OUT_XLSX)
        path = OUT_XLSX
    except OSError:
        alt = OUT / "ОПД_ДЗ_Таблицы_new.xlsx"
        if alt.exists():
            alt.unlink()
        tmp.replace(alt)
        path = alt
        print("WARN: файл занят в Excel — записан:", path)
    # убрать старый дубликат _new, если основной записался
    if path == OUT_XLSX:
        stale = OUT / "ОПД_ДЗ_Таблицы_new.xlsx"
        if stale.exists():
            try:
                stale.unlink()
            except OSError:
                pass
    print("xlsx:", path)
    return path


def build():
    wb = Workbook()

    # ——— Содержание ———
    ws = wb.active
    ws.title = "Содержание"
    ws["A1"] = "ОПД 2026 — таблицы домашнего задания"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:B1")
    rows_info = [
        ("Команда", "кузницабагов"),
        ("Проект", "Эхо-Рыцарь"),
        ("Жанр", "dungeon roguelike (Soul Knight-like), карта этажа 3×3"),
        ("Файл игры", "EchoKnight.c3p (ровно 50 событий Construct 3)"),
        ("Платформа", "Construct 3 free"),
        ("Пересборка", "python tools/export_office.py"),
    ]
    r = 3
    for k, v in rows_info:
        ws.cell(r, 1, k).font = BOLD
        ws.cell(r, 2, v).font = CELL_FONT
        r += 1
    r += 1
    ws.cell(r, 1, "Листы в этом файле").font = BOLD
    r += 1
    sheets = [
        ("Содержание", "Титул и оглавление"),
        ("ДЗ1_Самопрезентация", "Черновик формы"),
        ("ДЗ2_Риски", "Таблица рисков проекта"),
        ("ДЗ3_Аналитика", "Референсы + вывод"),
        ("ДЗ4_Концепт", "Концепт, карта, управление, ядра, сдача"),
        ("Чеклист", "Что сдать руками"),
    ]
    write_table(ws, r, ["Лист", "Что внутри"], sheets)
    autosize(ws, [28, 55])

    # ——— ДЗ1 ———
    ws = wb.create_sheet("ДЗ1_Самопрезентация")
    r = 1
    r = section_title(ws, r, "ДЗ1 — Самопрезентация", 2)
    r = para(ws, r, "Форма: https://forms.yandex.ru/cloud/6aa8916c068ff0fb048e9434 · Чат: https://vk.com/club241384912", 2, 30)
    write_table(ws, r, ["Поле", "Значение"], [
        ("ФИО", "_______________________"),
        ("Группа", "ИРИТ-РТФ, 1 курс"),
        ("Команда", "кузницабагов"),
        ("Роль", "разработчик на Construct 3"),
        ("О себе", "Учусь на 1 курсе УрФУ. В команде «кузницабагов» делаю роглайк «Эхо-Рыцарь» на Construct 3 (карта этажа 3×3, ядра эха, реликвии). Хочу довести игру до сдачи с полным циклом."),
        ("Навыки", "Construct 3, базовый JS/Python"),
        ("Контакт", "_______________________"),
    ], alt=True)
    autosize(ws, [22, 80])

    # ——— ДЗ2 ———
    ws = wb.create_sheet("ДЗ2_Риски")
    r = 1
    r = section_title(ws, r, "ДЗ2 — Риски проекта · Эхо-Рыцарь · кузницабагов", 4)
    r = para(ws, r, "YouGile: карточки ДЗ1–ДЗ4, билд EchoKnight.c3p, видео — у каждой ответственный, дедлайн, ссылка.", 4, 32)
    risks = [
        ("Нерабочая ссылка на видеопрезентацию", "−12,5 баллов от итоговой оценки", "высокая",
         "Перед сдачей открыть ссылку в инкогнито; копия на Я.Диске"),
        ("Проект только на ноутбуке разработчика", "−12,5 баллов; нельзя сдать; долг по ОПД", "средняя",
         "Общий доступ к EchoKnight.c3p (Диск/GitHub) всей команде"),
        ("Превышен лимит событий free Construct", "Нельзя править события; срыв сроков", "высокая",
         "Аккаунт с почтой (лимит 50); держать ровно 50; сборка node tools/build_echo.js"),
        ("Баг карты / дверей / кузницы на проверке", "Демо выглядит сломанным", "высокая",
         "Чеклист: меню → реликвия → зачистка → двери → 2–3 комнаты → F кузница → B босс → R"),
        ("Разработчик пропал на дедлайн", "Нет билда", "средняя",
         "Минимум 2 человека умеют открыть .c3p; актуальная копия в облаке"),
        ("Нет / плохое видео геймплея", "Не принимают артефакт", "средняя",
         "Записать 60–120 сек: карта этажа + комнаты с врагами + кузница + Q"),
        ("Аналитика про Dota, а игра — роглайк", "Расхождение артефактов и билда", "средняя",
         "Все ДЗ и таблицы только про Эхо-Рыцаря / Soul Knight-лайки"),
        ("Потеря правок в браузерном Construct", "Откат прогресса", "средняя",
         "Версии EchoKnight_v01.c3p… + генератор build_echo.js"),
        ("Команду не находят в ведомости", "Баллов 0", "низкая",
         "Сверить ФИО в «Сбор команд»; скрин формы куратору"),
        ("Смена жанра в середине семестра", "Концепт и аналитика устаревают", "низкая",
         "Жанр зафиксирован: dungeon roguelike, карта 3×3, кузница ядер"),
    ]
    write_table(ws, r, ["Риск", "Последствие", "Вероятность", "Меры по управлению"], risks)
    autosize(ws, [36, 30, 14, 48])

    # ——— ДЗ3 ———
    ws = wb.create_sheet("ДЗ3_Аналитика")
    r = 1
    r = section_title(ws, r, "ДЗ3 — Аналитика референсов · Эхо-Рыцарь · кузницабагов", 3)
    analytics = [
        ("Soul Knight", "Карта этажа, комнаты, очистка → дверь, WASD + стрельба, dash",
         "Онлайн-кооп, магазин скинов, десятки героев"),
        ("Enter the Gungeon / Binding of Isaac", "Комнатный цикл, риск-награда, смерть = рестарт",
         "Сотни предметов, огромный procedural graph"),
        ("Warframe (ядра / моды)", "Идея слотов усиления оружия",
         "Открытый мир, крафт-экономика, мультиплеер"),
        ("Minecraft Dungeons", "Реликвии / артефакты как отдельный слот",
         "Кампания с миссиями, лут как у AAA"),
        ("Hades", "Читаемые баффы по стихиям",
         "Диалоги, meta-прогресс между забегами"),
        ("Vampire Survivors", "Простая петля «выжил → сильнее»",
         "Автоатака без прицела, орды на одном экране"),
    ]
    r = write_table(ws, r, ["Игра / референс", "Что берём", "Что не берём / упрощаем"], analytics)
    r += 1
    r = section_title(ws, r, "Вывод для «Эхо-Рыцаря»", 3)
    r = para(
        ws, r,
        "Короткий забег по карте этажа 3×3: комнаты с врагами → двери → кузница ядер (F) "
        "огонь/лёд/шок → реликвия Q → босс (B). Оригинальный крюк — кузница эха на клетке карты, "
        "а не копирование чужого лута. Бюджет Construct 3 — ровно 50 событий. Файл EchoKnight.c3p.",
        3, 70,
    )
    autosize(ws, [34, 48, 40])

    # ——— ДЗ4 ———
    ws = wb.create_sheet("ДЗ4_Концепт")
    r = 1
    r = section_title(ws, r, "ДЗ4 — Концепт-документ · Эхо-Рыцарь · кузницабагов", 4)
    r = para(ws, r, "Платформа: Construct 3 (≤50 событий) · Файл: EchoKnight.c3p · Команда: кузницабагов", 4, 28)

    r = section_title(ws, r, "1. Elevator pitch", 4)
    r = para(
        ws, r,
        "«Эхо-Рыцарь» — комнатный роглайк в духе Soul Knight: зачищай залы на карте этажа 3×3, "
        "собирай ядра эха в кузнице и добивай босса одним усиленным билдом за забег.",
        4, 48,
    )

    r = section_title(ws, r, "2. Жанр и референсы", 4)
    r = para(
        ws, r,
        "Dungeon roguelike / twin-stick shooter. Референсы: Soul Knight (карта и цикл комнат), "
        "Warframe/MCD (слоты усилений), Hades (читаемые стихии).",
        4, 40,
    )

    r = section_title(ws, r, "3. Карта этажа (как в Soul Knight)", 4)
    r = write_table(ws, r, ["Клетка", "Тип", "Что внутри", "Примечание"], [
        ("1,1 (центр)", "Старт", "2 слайма", "Начало забега"),
        ("остальные боевые", "Бой", "Слаймы + летучие мыши по углам/центру", "Двери после зачистки"),
        ("2,1", "Кузница F", "Ядра 1–3, апгрейды 4–6 за золото", "Выход — дверь на карту"),
        ("2,0", "Босс B", "Голем + адды", "Победа за этаж"),
    ])
    r += 1
    r = para(ws, r, "Миникарта: ◆ ты · ○ зачищено · B босс · F кузница · · ещё не открыто", 4, 28)

    r = section_title(ws, r, "4. Геймплейный цикл", 4)
    r = write_table(ws, r, ["№", "Шаг"], [
        ("1", "Меню → выбор реликвии → старт в центре карты"),
        ("2", "Комната с врагами → зачистка → открываются двери"),
        ("3", "Переход в соседнюю клетку карты"),
        ("4", "Кузница F: ядра 1–3, шоп 4–6 → выход дверью"),
        ("5", "Босс B → победа / смерть → R новый забег"),
    ])
    r += 1

    r = section_title(ws, r, "5. Управление", 4)
    r = write_table(ws, r, ["Клавиша / ввод", "Действие"], [
        ("WASD", "Движение"),
        ("ЛКМ", "Выстрел"),
        ("Пробел", "Рывок (dash)"),
        ("Q", "Реликвия"),
        ("1 / 2 / 3", "Ядро огня / льда / молнии в кузнице"),
        ("4 / 5 / 6", "Апгрейды за золото в кузнице"),
        ("Дверь", "Переход на соседнюю клетку карты"),
        ("P", "Пауза"),
        ("R", "В меню / новый забег"),
    ])
    r += 1

    r = section_title(ws, r, "6. Ядра эха (УТП)", 4)
    r = write_table(ws, r, ["Стихия", "Клавиша", "Эффект"], [
        ("Огонь", "1", "+урон, периодический дот ближайшему"),
        ("Лёд", "2", "замедление врагов рядом"),
        ("Молния / шок", "3", "быстрая стрельба + цепной урон при убийстве"),
    ])
    r += 1
    r = para(
        ws, r,
        "УТП: не клон Soul Knight, а кузница эха на карте этажа — игрок сам собирает стихию "
        "в середине забега и сразу чувствует смену темпа боя.",
        4, 48,
    )

    r = section_title(ws, r, "7. Объём сдачи", 4)
    write_table(ws, r, ["Артефакт", "Статус / что сделать"], [
        ("EchoKnight.c3p в облаке", "выложить ссылку команде"),
        ("Видео 60–120 сек", "карта + комнаты + кузница + Q"),
        ("Этот Excel ОПД_ДЗ_Таблицы.xlsx", "в форму / YouGile / диск"),
        ("YouGile", "карточки с ссылками"),
    ])
    autosize(ws, [22, 18, 48, 28])

    # ——— Чеклист ———
    ws = wb.create_sheet("Чеклист")
    r = 1
    r = section_title(ws, r, "Чеклист сдачи ОПД 2026 · кузницабагов", 3)
    write_table(ws, r, ["№", "Действие", "Готово"], [
        ("1", "ДЗ1 — форма Яндекс + вступление в VK курса", "☐"),
        ("2", "YouGile — доска команды, карточки с файлами", "☐"),
        ("3", "Облако — ссылка на EchoKnight.c3p", "☐"),
        ("4", "Видео 1–2 мин на YouTube (ФИО + ссылка на .c3p в описании)", "☐"),
        ("5", "Формы ДЗ2–ДЗ4 — прикрепить этот Excel / ссылку на диск", "☐"),
        ("6", "Проверка билда: карта 3×3, враги в комнатах, F, B, 50 событий", "☐"),
    ])
    autosize(ws, [6, 70, 10])

    return save_xlsx(wb)


if __name__ == "__main__":
    build()
